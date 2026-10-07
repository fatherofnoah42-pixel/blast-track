import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

cfg = json.load(open('cfg_default.json'))
OUT = sys.argv[1] if len(sys.argv) > 1 else '爆破跑道_配置表.xlsx'
F = 'Arial'
HDR = PatternFill('solid', fgColor='1A2437'); HF = Font(name=F, bold=True, color='FFFFFF')
KEYF = Font(name=F, italic=True, size=9, color='8A93A6'); KEYFILL = PatternFill('solid', fgColor='EEF0F4')
IDF = Font(name=F, color='5A6478'); INF = Font(name=F, color='0000FF'); TXF = Font(name=F, color='000000'); CALC = Font(name=F, color='000000', italic=True)
thin = Side(style='thin', color='D5D9E0'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(vertical='top', wrap_text=True)
wb = Workbook()

def table(ws, cols, rows, widths=None, freeze='C3'):
    """cols: list of (human header, key, kind) kind in id|num|text|calc ; rows: list of lists"""
    for j, (hd, key, kind) in enumerate(cols, 1):
        c = ws.cell(1, j, hd); c.font = HF; c.fill = HDR; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = BD
        k = ws.cell(2, j, key); k.font = KEYF; k.fill = KEYFILL; k.border = BD
    for i, row in enumerate(rows, 3):
        for j, v in enumerate(row, 1):
            kind = cols[j - 1][2]
            c = ws.cell(i, j, v); c.border = BD; c.alignment = WRAP
            c.font = {'id': IDF, 'num': INF, 'text': TXF, 'calc': CALC}[kind]
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[CL(j)].width = w
    ws.cell(2, 1).comment = Comment('第 2 行是给游戏读取的字段名，请不要修改。', 'config')

def L3(base, label):  # three language columns
    return [(f'{label}（中文）', f'{base}_zh', 'text'), (f'{label}（English）', f'{base}_en', 'text'), (f'{label}（Français）', f'{base}_fr', 'text')]
def V3(d): return [d['zh'], d['en'], d['fr']]

# ---------- 说明 ----------
ws = wb.active; ws.title = '说明'
lines = [
 ('爆破跑道 · 游戏配置表', True),
 ('', False),
 ('怎么用', True),
 ('1. 修改各工作表里的数值和名称（蓝色字是数值，黑色字是文字，灰色字是 id，不要改 id）。', False),
 ('2. 每个工作表的第 2 行是游戏读取用的字段名，请不要修改、删除或调换列顺序以外的内容；可以调换列顺序。', False),
 ('3. 在 Excel 中保存为 .xlsx；如果用 Google Sheets，先把本文件上传到 Google 云端硬盘并用 Google 表格打开，改完后选「文件 → 下载 → Microsoft Excel (.xlsx)」。', False),
 ('4. 打开游戏页面，点「导入配置表 (.xlsx)」选择文件。游戏会立即按新配置重新开局；有问题的单元格会列在导入提示里，并自动使用默认值。', False),
 ('5. 点「恢复默认」可以回到内置配置。导入的配置只保存在当前浏览器里，换设备需要重新导入。', False),
 ('', False),
 ('注意', True),
 ('· 说明文字不会随数值自动变化。改了数值（比如冷却秒数），记得同步改对应的说明文字（三种语言）。', False),
 ('· 名称、说明都有中文、English、Français 三列；某一列留空会使用原来的文字。', False),
 ('· 斜体黑字的列（如「累计花费」「本波总人数」）是自动计算的参考列，游戏不读取。', False),
 ('· 「波次」表的行数就是总波数，可以增删行；boss 列填 1 的波次会出现 Boss。', False),
 ('· 「Boss」「卡牌」表的 enabled 填 0 可以停用该 Boss 或卡牌。', False),
 ('· 「升级」表的 effect 列用简单语法描述效果，写法见「效果语法」工作表。', False),
 ('· 地图坐标以 360 × 600 的画布为准，左上角是 (0, 0)。一种跑道可以有多条路线，实现多进口、多出口。', False),
 ('· 局内 GST 只是关卡内货币，不和任何真实资产挂钩。', False),
 ('', False),
 ('工作表一览', True),
 ('全局：经济、抽卡权重、出怪节奏、各种机制倍率等全局参数', False),
 ('装置：6 种装置的名称、价格和基础数值', False),
 ('升级路线：每种装置 3 条路线的名称', False),
 ('升级：75 个升级的名称、价格、说明和效果', False),
 ('敌人：普通跑者、耐力跑者、Boss 的基础数值', False),
 ('波次：每一波的人数、体力、是否有 Boss、词缀数量', False),
 ('Boss：6 位 Boss 的名称、能力说明和能力参数 p1–p4', False),
 ('词缀：4 种对手词缀的名称和说明（数值在「全局」）', False),
 ('卡牌：神秘箱里的全部卡牌', False),
 ('跑道列表：6 种跑道（单进单出、双进双出、单进双出、双进单出、十字交叉、三进单出）的名称、启用开关和难度倍率', False),
 ('地图：每种跑道的路线点和装置空位坐标', False),
 ('美术资源：每张美术图的编号、图片地址、游戏内尺寸和锚点；留空就用游戏里的矢量占位图', False),
 ('界面文字：按钮、提示等界面文字的三语翻译', False),
 ('效果语法：「升级」表 effect 列的写法', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.font = Font(name=F, bold=b, size=14 if i == 1 else 11); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 120

# ---------- 全局 ----------
ws = wb.create_sheet('全局')
table(ws, [('参数', 'key', 'id'), ('数值', 'value', 'num'), ('说明', 'desc', 'text')],
      [[k, v, cfg['globalDesc'][k]] for k, v in cfg['global'].items()], [22, 12, 60], 'B3')

# ---------- 装置 ----------
ws = wb.create_sheet('装置')
cols = [('id', 'id', 'id')] + L3('name', '名称') + L3('short', '简称') + L3('desc', '说明') + [
 ('放置价格', 'cost', 'num'), ('效果', 'dmg', 'num'), ('出手间隔（秒）', 'rate', 'num'), ('射程（999=全图）', 'range', 'num'),
 ('范围半径', 'rad', 'num'), ('减速比例', 'slow', 'num'), ('湿滑持续（秒）', 'slowDur', 'num'), ('传播人数', 'chain', 'num'),
 ('传播衰减', 'decay', 'num'), ('对耐力跑者效果', 'armor', 'num'), ('对 Boss 效果', 'bossMul', 'num'), ('颜色', 'color', 'text')]
rows = [[t['id']] + V3(t['name']) + V3(t['short']) + V3(t['desc']) + [t[k] for k in ['cost','dmg','rate','range','rad','slow','slowDur','chain','decay','armor','bossMul','color']] for t in cfg['towers']]
table(ws, cols, rows, [9, 14, 16, 18, 10, 10, 12, 40, 40, 40] + [11] * 11 + [10])

# ---------- 升级路线 ----------
ws = wb.create_sheet('升级路线')
table(ws, [('装置 id', 'tower', 'id'), ('路线', 'path', 'id')] + L3('name', '路线名称'),
      [[p['tower'], p['path']] + V3(p['name']) for p in cfg['paths']], [10, 7, 16, 20, 24])

# ---------- 升级 ----------
ws = wb.create_sheet('升级')
cols = [('装置 id', 'tower', 'id'), ('路线', 'path', 'id'), ('等级', 'tier', 'id')] + L3('name', '名称') + [('价格', 'cost', 'num')] + L3('desc', '说明') + [('效果（见「效果语法」）', 'effect', 'num'), ('本路线累计花费', '', 'calc')]
rows = []
for p in cfg['paths']:
    for t in p['tiers']:
        rows.append([p['tower'], p['path'], t['tier']] + V3(t['name']) + [t['cost']] + V3(t['desc']) + [t['effect'], None])
table(ws, cols, rows, [9, 6, 6, 16, 20, 22, 9, 36, 40, 42, 28, 12], 'D3')
last = 2 + len(rows)
for i in range(3, last + 1):
    ws.cell(i, 12, f'=SUMIFS($G$3:$G${last},$A$3:$A${last},A{i},$B$3:$B${last},B{i},$C$3:$C${last},"<="&C{i})').font = CALC
    ws.cell(i, 12).border = BD

# ---------- 敌人 ----------
ws = wb.create_sheet('敌人')
table(ws, [('id', 'id', 'id')] + L3('name', '名称') + [('体型半径', 'radius', 'num'), ('移速', 'speed', 'num'), ('劝退 GST', 'gold', 'num'), ('计分权重', 'value', 'num'), ('漏过扣能量', 'leak', 'num')],
      [[e['id']] + V3(e['name']) + [e['radius'], e['speed'], e['gold'], e['value'], e['leak']] for e in cfg['enemies']], [8, 14, 20, 22, 10, 10, 10, 10, 12])
ws.cell(3 + len(cfg['enemies']) + 1, 1, '※ 体力按波次配置，见「波次」表。').font = Font(name=F, italic=True, color='5A6478')

# ---------- 波次 ----------
ws = wb.create_sheet('波次')
cols = [('波次', 'wave', 'id'), ('慢跑者人数', 'normals', 'num'), ('耐力跑者人数', 'armored', 'num'), ('慢跑者体力', 'normHp', 'num'), ('耐力跑者体力', 'armHp', 'num'),
        ('是否 Boss 波（1/0）', 'boss', 'num'), ('Boss 体力', 'bossHp', 'num'), ('词缀数量', 'affixes', 'num'), ('本波总人数', '', 'calc')]
rows = [[w['wave'], w['normals'], w['armored'], w['normHp'], w['armHp'], w['boss'], w['bossHp'], w['affixes'], None] for w in cfg['waves']]
table(ws, cols, rows, [8, 12, 13, 12, 13, 16, 11, 10, 12], 'B3')
for i in range(3, 3 + len(rows)):
    ws.cell(i, 9, f'=B{i}+C{i}+F{i}').font = CALC; ws.cell(i, 9).border = BD
dv = DataValidation(type='list', formula1='"0,1"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'F3:F{200}')
ws.cell(3 + len(rows) + 1, 1, '※ 行数就是总波数，可以增删行；精英小队的词缀数量见「全局」eliteAffixes。').font = Font(name=F, italic=True, color='5A6478')

# ---------- Boss ----------
ws = wb.create_sheet('Boss')
cols = [('id', 'id', 'id'), ('启用（1/0）', 'enabled', 'num'), ('名字', 'name', 'text')] + L3('title', '称号') + L3('ability', '能力说明') + [
        ('p1', 'p1', 'num'), ('p2', 'p2', 'num'), ('p3', 'p3', 'num'), ('p4', 'p4', 'num'), ('p1–p4 含义', '', 'calc')]
rows = []
for b in cfg['bosses']:
    meaning = '；'.join(f'p{i+1} = {m}' for i, m in enumerate(cfg['bossParams'][b['id']]) if m)
    rows.append([b['id'], b['enabled'], b['name']] + V3(b['title']) + V3(b['ability']) + b['p'] + [meaning])
table(ws, cols, rows, [8, 10, 12, 12, 14, 16, 40, 44, 48, 8, 8, 8, 8, 46], 'D3')
dv = DataValidation(type='list', formula1='"0,1"'); ws.add_data_validation(dv); dv.add('B3:B20')

# ---------- 词缀 ----------
ws = wb.create_sheet('词缀')
table(ws, [('id', 'id', 'id')] + L3('name', '名称') + L3('desc', '说明'),
      [[a['id']] + V3(a['name']) + V3(a['desc']) for a in cfg['affixes']], [9, 12, 16, 18, 26, 36, 36])
ws.cell(3 + len(cfg['affixes']) + 1, 1, '※ 词缀数值（冲刺倍率、补给回复、能量盾比例、跟跑者体力）在「全局」表。').font = Font(name=F, italic=True, color='5A6478')

# ---------- 卡牌 ----------
ws = wb.create_sheet('卡牌')
cols = [('id', 'id', 'id'), ('启用（1/0）', 'enabled', 'num'), ('类别', 'tag', 'id'), ('稀有度（C/R/L）', 'rar', 'num'), ('可重复（1/0）', 'stack', 'num')] + L3('name', '名称') + L3('desc', '说明') + [
        ('数值 v1', 'v1', 'num'), ('数值 v2', 'v2', 'num'), ('v1 / v2 含义', '', 'calc')]
rows = [[c['id'], c['enabled'], c['tag'], c['rar'], c['stack']] + (V3(c['name']) if c['name'] else ['（自动）'] * 3) + (V3(c['desc']) if c['desc'] else ['（自动）'] * 3) + [c['v1'], c['v2'], c['note']] for c in cfg['cards']]
table(ws, cols, rows, [10, 10, 7, 12, 11, 14, 18, 20, 34, 38, 40, 9, 9, 30], 'F3')
for rng, f in [('B3:B60', '"0,1"'), ('D3:D60', '"C,R,L"'), ('E3:E60', '"0,1"')]:
    dv = DataValidation(type='list', formula1=f); ws.add_data_validation(dv); dv.add(rng)
ws.cell(3 + len(rows) + 1, 1, '※ 徽章卡（m_ 开头）的名称和说明由装置简称自动生成，填写无效。稀有度权重在「全局」rarityC/R/L。').font = Font(name=F, italic=True, color='5A6478')

# ---------- 跑道列表 ----------
ws = wb.create_sheet('跑道列表')
cols = [('id', 'id', 'id'), ('启用（1/0）', 'enabled', 'num')] + L3('name', '名称') + L3('desc', '说明') + [('对手移速倍率', 'spdMult', 'num'), ('对手体力倍率', 'hpMult', 'num'), ('入口 / 出口 / 路线数', '', 'calc')]
rows = []
for m in cfg['maps']:
    ins = len({tuple(r[0]) for r in m['routes']}); outs = len({tuple(r[-1]) for r in m['routes']})
    rows.append([m['id'], m['enabled']] + V3(m['name']) + V3(m['desc']) + [m['spdMult'], m['hpMult'], f'{ins} 进 {outs} 出 · {len(m["routes"])} 条路线'])
table(ws, cols, rows, [10, 10, 12, 14, 16, 40, 44, 48, 12, 12, 22], 'C3')
dv = DataValidation(type='list', formula1='"0,1"'); ws.add_data_validation(dv); dv.add('B3:B40')
ws.cell(3 + len(rows) + 1, 1, '※ 每局开始时从启用的跑道里随机选一条（游戏里也可以手动指定）。新增跑道：在这里加一行 id，再到「地图」表填坐标。').font = Font(name=F, italic=True, color='5A6478')

# ---------- 地图 ----------
ws = wb.create_sheet('地图')
rows = []
for m in cfg['maps']:
    for ri, r in enumerate(m['routes'], 1):
        for i, (x, y) in enumerate(r, 1): rows.append([m['id'], 'path', ri, i, x, y])
    for i, (x, y) in enumerate(m['pads'], 1): rows.append([m['id'], 'pad', '', i, x, y])
table(ws, [('跑道 id', 'map', 'id'), ('类型（path=路线点，pad=空位）', 'type', 'id'), ('路线编号', 'route', 'num'), ('顺序', 'order', 'num'), ('x', 'x', 'num'), ('y', 'y', 'num')], rows, [10, 28, 10, 8, 8, 8], 'A3')
ws.cell(3 + len(rows) + 1, 1, '※ 画布 360 × 600，左上角是 (0, 0)。同一跑道可以有多条路线（route 1、2、3…），每条路线按顺序连成折线，建议只画水平或竖直线段；入口和出口可以在画布外。多条路线共用的路段要用相同坐标，看起来就是一条路。跑者按路线轮流出发。空位离跑道中心线至少 36，空位之间至少 34。').font = Font(name=F, italic=True, color='5A6478')


# ---------- 美术资源 ----------
ws = wb.create_sheet('美术资源')
cols = [('资源编号（勿改）', 'key', 'id'), ('类别', '', 'calc'), ('说明', '', 'calc'), ('优先级', '', 'calc'), ('图片地址（https:// 链接，留空 = 用占位图）', 'src', 'text'),
        ('游戏内宽', 'w', 'num'), ('游戏内高', 'h', 'num'), ('锚点 X（0–1）', 'ax', 'num'), ('锚点 Y（0–1）', 'ay', 'num'), ('帧数', 'frames', 'num'), ('帧率', 'fps', 'num'), ('盖住跑道（1/0，仅地图）', 'cover', 'num'),
        ('建议出图尺寸（3 倍，px）', '', 'calc')]
rows = [[a['key'], a['kind'], a['desc'], a['prio'], a['src'], a['w'], a['h'], a['ax'], a['ay'], a['frames'], a['fps'], a['cover'], None] for a in cfg['art']]
table(ws, cols, rows, [20, 10, 52, 11, 46, 9, 9, 11, 11, 7, 7, 12, 22], 'B3')
for i in range(3, 3 + len(rows)):
    c = ws.cell(i, 13, f'=IF(J{i}>1,J{i}*F{i}*3&" × "&G{i}*3&"（"&J{i}&" 帧横排）",F{i}*3&" × "&G{i}*3)'); c.font = CALC; c.border = BD
dv = DataValidation(type='list', formula1='"0,1"'); ws.add_data_validation(dv); dv.add(f'L3:L{2 + len(rows)}')
notes = [
 '※ 图片地址填公开可访问的 https:// 链接（PNG / WebP，透明背景）；也可以不填，直接在游戏页面点「导入美术图片」，把文件名改成资源编号（如 tower.dart.png）一次选多张。',
 '※ 游戏内宽高以 360 × 600 的画布为单位；出图按 3 倍像素（见最后一列），游戏会缩放到这里填的尺寸。',
 '※ 锚点是图片里对齐游戏坐标的那一点：0,0 = 左上角，0.5,1 = 底边中点（跑者、装置用脚底中点）。',
 '※ 帧数 > 1 时，图片是横向排开的序列帧（每帧等宽），按帧率循环播放。跑者、Boss 统一画成「侧面向右跑」，向左时游戏自动水平翻转。',
 '※ 装置进化图（tower.xxx.p1–p3）在该路线升到 3 级以上并成为主路线时替换基础图；没给就继续用基础图。',
 '※ 优先级：P1 样板关（先画这些就能完整玩一张图）→ P2 → P3。',
]
for k, t in enumerate(notes):
    ws.cell(3 + len(rows) + 1 + k, 1, t).font = Font(name=F, italic=True, color='5A6478')

# ---------- 界面文字 ----------
ws = wb.create_sheet('界面文字')
table(ws, [('原文（id，勿改）', 'key', 'id'), ('中文', 'zh', 'text'), ('English', 'en', 'text'), ('Français', 'fr', 'text')],
      [[t['key'], t['zh'], t['en'], t['fr']] for t in cfg['texts']], [44, 44, 50, 54], 'B3')
ws.cell(3 + len(cfg['texts']) + 1, 1, '※ 花括号里的 {n}、{x} 等是占位符，翻译时要原样保留。').font = Font(name=F, italic=True, color='5A6478')

# ---------- 效果语法 ----------
ws = wb.create_sheet('效果语法')
syntax = [
 ('写法', '多个效果用分号隔开，例如：dmg+=2; armor=1'),
 ('+=', '增加：dmg+=1'), ('*=', '乘以：rate*=0.8（出手间隔 ×0.8，即变快 20%）'), ('=', '设为：stun=0.6'), ('>=', '至少为：armor>=0.6'),
 ('', ''), ('属性', '含义'),
 ('dmg', '效果（伤害）'), ('rate', '出手间隔（秒，越小越快）'), ('range', '射程'), ('rad', '香蕉皮范围半径'), ('slow', '减速比例'), ('slowDur', '湿滑持续（秒）'),
 ('chain', '音响传播人数'), ('decay', '音响每传一人的效果保留比例'), ('multi', '易拉罐每次砸几人'), ('crit', '额外暴击率'), ('critMul', '暴击倍率'),
 ('armor', '对耐力跑者的效果比例'), ('bossMul', '对 Boss 的效果倍率'), ('starts', '音响同时放几首'), ('shieldMul', '对能量盾的效果倍率'), ('bombs', '一次扔几片香蕉皮'),
 ('execChance', '命中直接绊倒退赛的概率（非 Boss）'), ('mark', '汽水黏鞋持续秒数'), ('auraRate', '光环：其他装置出手间隔倍率'), ('auraR', '光环半径（不填则等于射程）'),
 ('auraDmg', '光环：其他装置效果倍率'), ('stun', '眩晕秒数'), ('shrapnel', '香蕉碎片块数'), ('cluster', '小香蕉分裂层数'), ('burnDur', '辣椒持续秒数'),
 ('burnPct', '辣椒每秒效果比例'), ('firePatch', '辣油留存秒数'), ('inferno', '1 = 辣到退赛时溅射'), ('freeze', '冻结秒数'), ('bossSlowFull', '1 = Boss 吃满减速'),
 ('iceAge', '人工降雪间隔秒数'), ('vulnSlow', '被淋湿后额外受到的效果比例'), ('vulnFreeze', '被冻住后额外受到的效果比例'), ('strip', '泡水鞋持续秒数'),
 ('exec25', '1 = 体力低于 25% 直接退赛'), ('decap', '1 = 开启神秘箱轰炸'), ('ricochet', '鞋盒弹跳次数'), ('gpw', '每波结束额外 GST'),
 ('globalDmg', '全场装置效果加成'), ('overN', '每第几次触发重低音'), ('overMul', '重低音效果倍率'), ('staticDps', '魔性旋律每秒效果'),
]
table(ws, [('写法 / 属性', '', 'id'), ('说明', '', 'text')], [list(r) for r in syntax], [16, 70], 'A3')

wb.save(OUT)
print('saved', OUT)
